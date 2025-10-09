import uuid
from datetime import datetime
from typing import List, Any, Dict
import logging
import zipfile
import openpyxl
from sqlalchemy.orm import Session
from app.database import SessionLocal
from app.models import BoardCertificate, Application, BoardCertificateFile
from app.services.utils.excel_utils import parse_excel_date, normalize_keys, get_first_value
logger = logging.getLogger(__name__)


def find_header_row(worksheet) -> int:
    max_scan = min(8, worksheet.max_row)
    for row_idx in range(1, max_scan + 1):
        cell_values = [(str(cell.value) or "").strip().lower() for cell in worksheet[row_idx]]
        if any("npi" in v for v in cell_values) or any(v.startswith("board_abms_") for v in cell_values) or any("provider_name" in v for v in cell_values):
            return row_idx
    return 1


def load_board_certificate_rows(file_path: str) -> List[Dict[str, Any]]:
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    worksheet = workbook.worksheets[0]
    header_row_idx = find_header_row(worksheet)

    headers_by_index: Dict[int, str] = {}
    for col_idx, cell in enumerate(worksheet[header_row_idx], start=1):
        headers_by_index[col_idx] = (str(cell.value) or "").strip()

    records: List[Dict[str, Any]] = []
    for row_idx in range(header_row_idx + 1, worksheet.max_row + 1):
        record: Dict[str, Any] = {}
        is_empty = True
        for col_idx, header in headers_by_index.items():
            value = worksheet.cell(row=row_idx, column=col_idx).value
            record[header] = value
            if value not in (None, ""):
                is_empty = False
        if not is_empty:
            records.append(record)
    return records


def ingest_board_certificate_file(file_path: str) -> bool:
    db: Session = SessionLocal()
    try:
        records = load_board_certificate_rows(file_path)
        inserted_count = 0

        # Map NPI -> application_id
        npi_to_application_id = {}
        for application in db.query(Application).all():
            if application.npi and application.id:
                npi_to_application_id[str(application.npi).strip()] = application.id

        print(f"Loaded {len(records)} rows from Excel")
        print(f"Mapped {len(npi_to_application_id)} NPIs from applications table")

        for record in records:
            normalized_record = normalize_keys(record)
            npi = get_first_value(normalized_record, ["npi", "providernpi", "provider_npi"])
            if not npi:
                continue

            npi_str = str(int(npi)) if isinstance(npi, (int, float)) else str(npi).strip()
            application_id = npi_to_application_id.get(npi_str)
            if not application_id:
                print(f"⚠️ Skipping NPI {npi_str}: not found in applications table")
                continue

            existing = db.query(BoardCertificate).filter_by(application_id=application_id).first()
            if existing:
                continue

            board_certificate = BoardCertificate(
                boardcertificate_key=str(uuid.uuid4()),   # cast UUID to string
                application_id=application_id,
                abmsuid=get_first_value(normalized_record, ["boardabmsuid", "abmsuid"]),
                abms_name=get_first_value(normalized_record, ["providername", "provider_name", "abmsname"]),
                abms_dob=parse_excel_date(get_first_value(normalized_record, ["boardabmsdob", "dob"])),
                abms_education=get_first_value(normalized_record, ["boardabmseducation"]),
                abms_address=get_first_value(normalized_record, ["boardabmsaddress"]),
                abms_certification_board=get_first_value(normalized_record, ["boardabmscertificationboard"]),
                abms_certification_type=get_first_value(normalized_record, ["boardabmscertificationtype"]),
                abms_status=get_first_value(normalized_record, ["boardabmsstatus"]),
                abms_duration=get_first_value(normalized_record, ["boardabmsduration"]),
                abms_occurrence=get_first_value(normalized_record, ["boardabmsoccurrence"]),
                abms_start_date=parse_excel_date(get_first_value(normalized_record, ["boardabmsstartdate", "boardabmsstart_date"])),
                abms_end_date=parse_excel_date(get_first_value(normalized_record, ["boardabmsenddate", "boardabmsend_date"])),
                abms_reverification_date=parse_excel_date(get_first_value(normalized_record, ["boardabmsreverificationdate"])),
                abms_participating_in_moc=get_first_value(normalized_record, ["boardabmsparticipatinginmoc"]),
            )

            db.add(board_certificate)
            inserted_count += 1

        db.commit()
        print(f"BoardCertificate ingest complete. Inserted {inserted_count} rows.")

        # Verify with a new session
        with SessionLocal() as verify_db:
            result = verify_db.query(BoardCertificate).count()
            print(f" BoardCertificate count in DB: {result}")

        return True

    except Exception as e:
        db.rollback()
        print(f" Ingestion failed: {e}")
        return False
    finally:
        db.close()

def ingest_boardcertificate_files(zip_path):
            session = SessionLocal()
            try:
                # Build NPI -> boardcertificate_key mapping
                npi_to_boardcertificate_key = {}
                board_certificates = session.query(BoardCertificate).all()
                for bc in board_certificates:
                    app = session.query(Application).filter_by(id=bc.application_id).first()
                    if app and app.npi:
                        npi_to_boardcertificate_key[str(app.npi).strip()] = bc.boardcertificate_key

                print(f"Total mapped NPIs: {len(npi_to_boardcertificate_key)}")

                with zipfile.ZipFile(zip_path, "r") as zf:
                    for file_info in zf.infolist():
                        filename = file_info.filename
                        # Skip directories
                        if file_info.is_dir():
                            continue
                        base_name = filename.split("/")[-1]
                        npi_candidate = base_name.split(".")[0]

                        boardcertificate_key = npi_to_boardcertificate_key.get(npi_candidate)
                        if boardcertificate_key:
                            # Check if file already exists for this boardcertificate_key and filename
                            existing_file = session.query(BoardCertificateFile).filter_by(
                                boardcertificate_key=boardcertificate_key,
                                filename=base_name
                            ).first()
                            if existing_file:
                                print(f"File for {base_name} and boardcertificate_key {boardcertificate_key} already exists, skipping.")
                                continue
                            try:
                                file_data = zf.read(file_info.filename)
                                if file_data:
                                    print(f"File data found for {filename}, adding to DB.")
                                    board_file = BoardCertificateFile(
                                        boardcertificate_key=boardcertificate_key,
                                        filename=base_name,
                                        file_type=base_name.split(".")[-1],
                                        file_data=file_data,
                                    )
                                    session.add(board_file)
                                else:
                                    print(f"File data is empty for {filename}, skipping DB insert.")
                            except KeyError:
                                print(f"File {filename} not found in zip archive.")
                session.commit()
                print("BoardCertificate files ingested.")
            except Exception as e:
                session.rollback()
                print(f"Error: {e}")
            finally:
                session.close()