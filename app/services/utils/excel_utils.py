from datetime import datetime
from typing import Dict, Any, List

import openpyxl


def parse_excel_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    except Exception:
        return None


def find_header_row(worksheet) -> int:
    max_scan = min(8, worksheet.max_row)
    for row_idx in range(1, max_scan + 1):
        cell_values = [(str(cell.value) or "").strip().lower() for cell in worksheet[row_idx]]
        if any("npi" in v for v in cell_values) or any(v.startswith("board_abms_") for v in cell_values) or any("provider_name" in v for v in cell_values):
            return row_idx
    return 1


def normalize_keys(record: Dict[str, Any]) -> Dict[str, Any]:
    out = {}
    import re
    for key, value in record.items():
        normalized_key = re.sub(r"[^a-z0-9]", "", (str(key) or "").lower())
        out[normalized_key] = value
    return out



def get_first_value(normalized_record: Dict[str, Any], candidates: List[str]):
    for candidate in candidates:
        if candidate in normalized_record and normalized_record[candidate] not in (None, ""):
            return normalized_record[candidate]
    return None


def load_excel_rows(path: str) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(path, data_only=True)

    ws = wb.worksheets[0]

    header_row = find_header_row(ws)
    headers_by_idx = {j: (str(c.value) or "").strip() for j, c in enumerate(ws[header_row], start=1)}

    rows = []
    for i in range(header_row + 1, ws.max_row + 1):
        rec = {}
        empty = True
        for j, key in headers_by_idx.items():
            val = ws.cell(row=i, column=j).value
            rec[key] = val
            if val not in (None, ""):
                empty = False
        if not empty:
            rows.append(rec)
    return rows

def parse_cell_date(val):
    """Ensure dates are stored as date objects for DB."""
    return parse_excel_date(val)